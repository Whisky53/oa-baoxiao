#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os, sys, glob, re, time
from pathlib import Path
from pypdf import PdfReader
from playwright.sync_api import sync_playwright
try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None
sys.path.insert(0, str(Path(__file__).parent))
from rail_invoice_parse import parse_rail_pdf
from fill_oa_auto import (load_config, detect_chrome, arg, js_set, js_select2,
 select2_click_pick, select2_search_pick, relation_choose_pick, apply_post_rules, save_draft)

SKILL_DIR=str(Path(__file__).parent.parent)
CHROME=os.environ.get('CHROME') or detect_chrome()

def didi_parse(path):
    text='\n'.join((p.extract_text() or '') for p in PdfReader(path).pages)
    if '滴滴' not in text and 'didi' not in text.lower(): raise ValueError('非滴滴文件')
    invno=re.search(r'发票号码\s*[:：]?\s*(\d{8,})',text)
    total=re.search(r'价\s*税\s*合\s*计.*?([0-9]+(?:\.[0-9]{1,2})?)\s*¥',text,re.S)
    date_range=re.search(r'行程起止日期\s*[:：]?\s*(\d{4})\s*[年./-]\s*(\d{1,2})\s*[月./-]\s*(\d{1,2})\s*(?:日)?\s*(?:至|到|[-—~～]|\n)\s*(?:(\d{4})\s*[年./-]\s*)?(\d{1,2})\s*[月./-]\s*(\d{1,2})\s*(?:日)?',text,re.S)
    if date_range:
        y1,m1,d1,y2,m2,d2=date_range.groups(); y2=y2 or y1
        start_date=f'{y1}-{int(m1):02d}-{int(d1):02d}'; end_date=f'{y2}-{int(m2):02d}-{int(d2):02d}'
    else:
        single_date=re.search(r'(?:出行日期|乘车日期)\s*[:：]?\s*(\d{4})\s*[年./-]\s*(\d{1,2})\s*[月./-]\s*(\d{1,2})',text)
        start_date=f'{single_date.group(1)}-{int(single_date.group(2)):02d}-{int(single_date.group(3)):02d}' if single_date else ''
        end_date=start_date
    route=re.search(r'广州\s*市.*?\n.*?公司',text,re.S)
    is_invoice=('电子发票' in text or '发票号码' in text)
    return {'kind':'didi_invoice' if is_invoice else 'didi_trip','发票号码':invno.group(1) if invno else '', '价税合计':float(total.group(1)) if total else 0.0, '乘车日期':start_date, '开始日期':start_date, '结束日期':end_date, '出发站':'广州','到达站':'东莞','金额':float(total.group(1)) if total else 0.0, '_file':path}

def hotel_parse(path):
    text='\n'.join((p.extract_text() or '') for p in PdfReader(path).pages)
    if not any(x in text for x in ('住宿','酒店','代订住宿')): raise ValueError('非酒店发票')
    invno=re.search(r'(?<!\d)(\d{20})(?!\d)',text)
    date=re.search(r'(\d{4})年(\d{1,2})月(\d{1,2})日',text)
    total=re.search(r'¥\s*([0-9]+(?:\.[0-9]{1,2})?)',text)
    taxline=re.search(r'金\s*额\s*税率.*?\n.*?([0-9]+(?:\.[0-9]{1,2})?)\s+6%\s+([0-9]+(?:\.[0-9]{1,2})?)',text,re.S)
    d=f'{date.group(1)}-{int(date.group(2)):02d}-{int(date.group(3)):02d}' if date else ''
    return {'kind':'hotel','发票号码':invno.group(1) if invno else '', '价税合计':float(total.group(1)) if total else 0.0, '不含税金额':float(taxline.group(1)) if taxline else 0.0, '税额':float(taxline.group(2)) if taxline else 0.0, '乘车日期':d, '出发站':'','到达站':'','_file':path}

def flight_parse(path):
    text='\n'.join((p.extract_text() or '') for p in PdfReader(path).pages)
    if '电子发票' not in text or ('航空' not in text and '客票' not in text and '南方航空' not in text): raise ValueError('非机票发票')
    invno=re.search(r'发票号码\s*[:：]?\s*(?:\n|\s)*?(\d{20})',text)
    if not invno:
        tail=text[text.find('发票号码'):] if '发票号码' in text else text
        invno=re.search(r'(?<!\d)(\d{20})(?!\d)',tail)
    total=re.search(r'价税合计.*?¥\s*([0-9]+(?:\.[0-9]{1,2})?)',text,re.S) or re.search(r'¥\s*([0-9]+(?:\.[0-9]{1,2})?)',text)
    date=re.search(r'(\d{4})年(\d{1,2})月(\d{1,2})日',text)
    route=re.search(r'(广州)\s*[-—至]\s*(上海)',text)
    flight=re.search(r'([A-Z]{2}\d{3,4})',text)
    travel_date=re.search(r'(\d{1,2})月(\d{1,2})日',text)
    return {'kind':'flight','发票号码':invno.group(1) if invno else '', '价税合计':float(total.group(1)) if total else 0.0, '乘车日期':f'2026-{int(travel_date.group(1)):02d}-{int(travel_date.group(2)):02d}' if travel_date else '', '出发站':route.group(1) if route else '广州','到达站':route.group(2) if route else '上海','航班号':flight.group(1) if flight else '', '_file':path}

def main():
    cfg=load_config(); directory=sys.argv[1] if len(sys.argv)>1 and not sys.argv[1].startswith('-') else cfg.get('invoice_dir')
    if not directory: raise SystemExit('缺少票据目录')
    project=arg('--project') or cfg.get('project')
    if not project: raise SystemExit('缺少项目，请使用 --project')
    headless='--headless' in sys.argv; no_save='--no-save' in sys.argv
    state=os.environ.get('OA_STATE_FILE') or str(Path(SKILL_DIR)/'data/oa_state.json')
    files=sorted(glob.glob(os.path.join(directory,'*')))
    files=[f for f in files if os.path.isfile(f) and not os.path.basename(f).startswith('.')]
    if not files: raise SystemExit('票据目录无文件: '+directory)
    invoices=[]; trips=[]; extras=[]; seen=set(); detected_types=set()
    classified=[]
    # 恢复原逻辑：目录内每个文件都进入处理，不做票种预检筛选。
    for f in files:
        name=os.path.basename(f)
        kind='pdf' if f.lower().endswith('.pdf') else 'extra'
        classified.append((f,kind)); detected_types.add(kind)
    print('[全量处理]', len(files), '个文件')
    for f, kind in classified:
        try:
            if kind=='pdf':
                base=os.path.basename(f)
                if '自驾发票' in base:
                    inv={'kind':'drive','发票号码':'','价税合计':0.0,'乘车日期':'','出发站':'','到达站':'','_file':f}
                    try:
                        _t='\n'.join((p.extract_text() or '') for p in PdfReader(f).pages)
                        _m=re.search(r'(?<!\d)(\d{20})(?!\d)',_t)
                        if _m: inv['发票号码']=_m.group(1)
                        _a=re.search(r'¥\s*([0-9]+(?:\.[0-9]{1,2})?)',_t)
                        if _a: inv['价税合计']=float(_a.group(1))
                    except Exception: pass
                    invoices.append(inv); print('[INFO] 自驾发票进入发票栏:',base,inv['发票号码']); continue
                if '登机凭证' in base:
                    extras.append(f); print('[INFO] 非发票转其它附件:',base); continue
                if '行程单' in base or '行程报销单' in base:
                    try:
                        t=didi_parse(f); t['_file']=f
                        trips.append(t)
                        print('[INFO] 滴滴行程单解析入 trips:',base)
                    except Exception as e:
                        extras.append(f); print('[INFO] 行程单按附件处理:',base,str(e)[:60])
                    continue
                try: inv=parse_rail_pdf(f); inv['kind']='train'
                except Exception:
                    try: inv=flight_parse(f)
                    except Exception:
                        try: inv=hotel_parse(f)
                        except Exception: inv=didi_parse(f)
                inv['_file']=f
                key=inv.get('发票号码') or ('file:'+os.path.basename(f))
                if key not in seen:
                    seen.add(key); invoices.append(inv)
            elif kind=='didi_invoice':
                inv=didi_parse(f)
                if inv['发票号码'] and inv['发票号码'] in seen: continue
                if inv['发票号码']: seen.add(inv['发票号码'])
                invoices.append(inv)
            elif kind=='didi_trip':
                trips.append(didi_parse(f))
            elif kind=='extra':
                extras.append(f)
            else:
                extras.append(f)
        except Exception as e:
            # 无法解析的文件不得强行计入发票。通用兜底：文件名含 20 位连续数字（如“发票号.pdf”命名）
            # 时按发票保留、号码取自文件名，金额等字段待人工补全；否则转其它附件。
            base_name = os.path.basename(f)
            m = re.search(r'(?<!\d)(\d{20})(?!\d)', base_name)
            if m:
                print('[INFO] 无法解析但文件名含20位发票号，按发票保留并待人工补字段:', base_name)
                invoices.append({'kind':'other_invoice','发票号码':m.group(1),'价税合计':0.0,'乘车日期':'','出发站':'','到达站':'','_file':f})
            else:
                print('[INFO] 非发票文件转其它附件:',base_name,str(e)[:80])
                extras.append(f)
    drive_rows=[]
    drive_xlsx=None
    xlsx=next((f for f in files if f.lower().endswith('.xlsx')), None)
    if any(inv.get('kind')=='drive' for inv in invoices) and xlsx and load_workbook:
        wb=load_workbook(xlsx, data_only=True); ws=wb.active
        headers=[str(c.value or '').strip() for c in ws[1]]; idx={h:i for i,h in enumerate(headers)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            date_value=row[idx.get('自驾日期',0)]; start=row[idx.get('出发地点',1)]; end=row[idx.get('到达地点',2)]; amount_value=row[idx.get('金额',6)]
            if date_value and start and end and amount_value is not None:
                date_text=date_value.strftime('%Y-%m-%d') if hasattr(date_value, 'strftime') else str(date_value).split(' ')[0]
                drive_rows.append({'出发站':str(start).strip(),'到达站':str(end).strip(),'乘车日期':date_text,'公里数':row[idx.get('公里数',3)],'金额':float(amount_value or 0),'费用类型':'自驾','原因':row[idx.get('自驾原因',4)] or ''})
        drive=next(inv for inv in invoices if inv.get('kind')=='drive'); drive['drive_rows']=drive_rows
        drive_xlsx=xlsx
        # 自驾行程表 Excel 只上传到“自驾行程表”专用控件，不允许同时出现在其它附件栏
        if drive_xlsx in extras: extras.remove(drive_xlsx)
        print('[自驾Excel]',len(drive_rows),'条明细；行程表仅上传自驾专用控件，不进入其它附件')
    for inv in invoices:
        if inv['kind']=='didi_invoice' and trips:
            t=min(trips,key=lambda x:abs(x.get('金额',0)-inv.get('价税合计',0)))
            inv.update({'出发站':t['出发站'],'到达站':t['到达站'],'乘车日期':t['乘车日期'],'开始日期':t.get('开始日期', t.get('乘车日期','')),'结束日期':t.get('结束日期', t.get('乘车日期','')),'trip_file':t['_file'],'费用类型':'出租车/网约车费/滴滴'})
    if not invoices:
        raise SystemExit('发票数量校验失败：未识别到可确认发票；不暂存')
    print('数量校验通过：目录文件',len(files),'个，实际识别发票',len(invoices),'张，其它附件',len(extras),'个')
    with sync_playwright() as p:
        cdp=os.environ.get('OA_CDP_URL')
        using_cdp=False
        if cdp:
            try:
                browser=p.chromium.connect_over_cdp(cdp)
                using_cdp=True
            except Exception as e:
                print('[WARN] CDP连接失败，改用已保存OA会话启动可视化Chrome:', str(e).splitlines()[0])
                browser=p.chromium.launch(executable_path=CHROME,headless=False,slow_mo=30)
        else:
            browser=p.chromium.launch(executable_path=CHROME,headless=headless,slow_mo=30)
        ctx=browser.contexts[0] if using_cdp and browser.contexts else browser.new_context(storage_state=state,viewport={'width':1600,'height':1000})
        page=ctx.pages[-1] if using_cdp and ctx.pages else ctx.new_page()
        page.goto(cfg['oa']['reimburse_new_url'],timeout=60000); page.wait_for_load_state('domcontentloaded'); time.sleep(5)
        page = next((pg for pg in ctx.pages if 'kmReviewMain' in pg.url), page)
        page.locator('input[name="__landray_filefd_invoice_att"]').wait_for(state='attached', timeout=60000)
        js_set(page,'fd_377ebcbaabff4c',str(len(invoices)))
        select2_click_pick(page,'fd_company_code','__comp',keyword=cfg.get('company',''))
        js_select2(page,'fd_paper_type','电子票'); js_select2(page,'fd_377ebcdd332a46','是'); js_set(page,'fd_377f75c5120d22','出差客户现场'); js_select2(page,'fd_reimbursement_money_type','人民币')
        invoice_files=[x['_file'] for x in invoices]
        other_files=[]
        for inv in invoices:
            if inv.get('trip_file'): other_files.append(inv['trip_file'])
        for t in trips:  # 未匹配到发票的行程单也作为附件上传
            if t.get('_file') and t['_file'] not in other_files:
                other_files.append(t['_file'])
        other_files.extend(extras)
        page.set_input_files('input[name="__landray_filefd_invoice_att"]',invoice_files); time.sleep(6)
        if other_files:
            page.locator('input[name="__landray_filefd_other_att"]').set_input_files(other_files)
            time.sleep(3)
        if drive_xlsx:
            page.locator('input[name="__landray_filefd_self_driving_file"]').set_input_files(drive_xlsx)
            time.sleep(4)
            print('[自驾行程表] Excel 已上传到自驾行程表专用控件:', os.path.basename(drive_xlsx))
        try:
            select2_click_pick(page,'fd_378057c5573e88.0.fd_377fb4b73ae7d2','__trip',first=True)
        except Exception:
            page.evaluate("() => document.querySelectorAll('.lui_dialog_mask,.select2-drop-mask').forEach(x=>x.remove())")
            time.sleep(1)
        # 发票上传后由 OA 自动创建明细行（数量=发票数，顺序=上传顺序）。
        # 自驾发票：发票金额>=报销总额即可；其自动行填 Excel 第 1 条，其余按 Excel 明细点击“添加行”追加。
        deadline=time.time()+45
        while time.time()<deadline:
            if page.locator('input[name^="fd_fee_item."][name$="fd_receipt_no"]').count() >= len(invoices): break
            time.sleep(1)
        def click_add_row(target_idx):
            # 精确定位费用明细子表(TABLE_DL_fd_fee_item)内的“添加行”按钮
            page.evaluate("""() => { const tbl=document.querySelector('#TABLE_DL_fd_fee_item'); const a=tbl&&[...tbl.querySelectorAll('a,span,input[type=button]')].find(x=>((x.innerText||x.value||'').trim()==='添加行')&&x.getAttribute('onclick')); if(a){a.click(); return true;} return false; }""")
            deadline=time.time()+15
            while time.time()<deadline:
                if page.locator(f'input[name^="fd_fee_item.{target_idx}."]').count()>0: break
                time.sleep(1)
        tasks=[]   # (row_index 或 'append', drive_detail或None, invoice)
        for i,inv in enumerate(invoices):
            if inv['kind']=='drive':
                dr_list=inv.get('drive_rows') or [{}]
                tasks.append((i, dr_list[0], inv))
                for extra in dr_list[1:]:
                    tasks.append(('append', extra, inv))
            else:
                tasks.append((i, None, inv))
        fill_count=len(invoices)
        for row_spec,dr,inv in tasks:
            if row_spec=='append':
                click_add_row(fill_count); row_index=fill_count; fill_count+=1
            else:
                row_index=row_spec
            pre=f'fd_fee_item.{row_index}.'
            if dr is not None:
                # 自驾明细：按 Excel 行填充，共用同一发票号码
                d_start=dr.get('乘车日期', inv.get('乘车日期','')); d_end=dr.get('乘车日期', d_start)
                start_date=str(d_start)[:10] if d_start else ''; end_date=str(d_end)[:10] if d_end else ''
                amount=dr.get('金额', 0); route_from=dr.get('出发站',''); route_to=dr.get('到达站','')
                inv_no=inv.get('发票号码',''); kw='自驾'; rc='自驾'; inv_kind='drive'
            else:
                kind=inv['kind']; typ='市内交通' if kind=='didi_invoice' else ('飞机' if kind=='flight' else ('住宿' if kind=='hotel' else ('其它' if kind=='other_invoice' else '火车')))
                kw=inv.get('费用类型',typ); rc=('出租车/网约车费/滴滴' if kind=='didi_invoice' else ('飞机' if kind=='flight' else ('住宿' if kind=='hotel' else ('其它' if kind=='other_invoice' else '火车/高铁'))))
                date=inv.get('乘车日期',''); start_date=inv.get('开始日期',date); end_date=inv.get('结束日期',date)
                amount=inv.get('价税合计',inv.get('票价(含税)',0)); route_from=inv.get('出发站',''); route_to=inv.get('到达站',''); inv_no=inv.get('发票号码',''); inv_kind=kind
            try:
                select2_search_pick(page,pre+'fd_erp_type',f'__erp{row_index}',keyword=kw,result_contains=rc)
            except Exception as e:
                print(f'[WARN] 行{row_index}费用类型交互异常，继续覆盖字段:',str(e)[:80])
            for k,v in {pre+'fd_377ebdd6f64cee':route_from,pre+'fd_377f76fb2a4068':route_to,pre+'fd_377ebe019348aa':start_date,pre+'fd_377ebe06621da2':end_date,pre+'fd_38da1066e1ec68':amount,pre+'fd_receipt_no':inv_no}.items(): js_set(page,k,v)
            js_select2(page,pre+'fd_Invoice_currency','人民币'); js_select2(page,pre+'fd_receipt_type','专票' if inv_kind=='train' else '普票'); relation_choose_pick(page,pre+'fd_project_name',project)
        apply_post_rules(page)
        save_result='未执行'
        if not no_save:
            save_result=save_draft(page)
            print('暂存结果:',save_result)
            time.sleep(5)
        print('已填入 OA，项目:',project,'；仅暂存:',not no_save,'；暂存操作:',save_result)
        if headless: browser.close()
        else: time.sleep(1200); browser.close()
if __name__=='__main__': main()
